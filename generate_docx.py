"""
generate_docx.py — แปลง markdown ธรรมดา (ที่ได้จากโหมดร่างเอกสาร/scrutinize ใน rag_worker.py)
เป็นไฟล์ .docx แบบโครงสร้างเปล่า (heading/paragraph/bullet) ไม่เลียนแบบ template บริษัทจริง
(ไม่มีตารางเซ็นสั่ง/ทวนสอบ, ประวัติแก้ไข, หัวกระดาษ) — ตามการตัดสินใจใน ADR-001 ข้อ 5

ผู้ใช้คัดลอกเนื้อหาจากไฟล์นี้ไปใส่ template จริงขององค์กรเองก่อนเสนออนุมัติ
"""
import io
import re

from docx import Document
from docx.shared import Pt, RGBColor

# จับทั้ง **bold** ของ markdown และมาร์ก [ต้องระบุ: ...] (ดู ADR-002) ในรอบเดียว เพื่อรักษาลำดับ
# ตำแหน่งข้อความให้ถูกต้องเวลาทั้งสองแบบอยู่ในบรรทัดเดียวกัน
_INLINE_RE = re.compile(r"\*\*(?P<bold>.+?)\*\*|(?P<marker>\[ต้องระบุ:[^\]]*\])")


def _add_markdown_line(doc: Document, line: str) -> None:
    """เพิ่มหนึ่งบรรทัด markdown ลงในเอกสาร — รองรับ heading (#/##/###), bullet (-/*),
    และ inline **bold** ภายในย่อหน้าธรรมดา/bullet เท่านั้น (ไม่รองรับ markdown ซับซ้อนกว่านี้)"""
    stripped = line.rstrip()

    if not stripped.strip():
        doc.add_paragraph("")
        return

    heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
    if heading_match:
        level = len(heading_match.group(1))
        text = heading_match.group(2).strip()
        doc.add_heading(text, level=level)
        return

    bullet_match = re.match(r"^[-*]\s+(.*)$", stripped)
    if bullet_match:
        para = doc.add_paragraph(style="List Bullet")
        _add_runs_with_bold(para, bullet_match.group(1))
        return

    para = doc.add_paragraph()
    _add_runs_with_bold(para, stripped)


def _add_runs_with_bold(paragraph, text: str) -> None:
    """แตก text เป็น run ปกติ/ตัวหนา ตามเครื่องหมาย **...** ของ markdown และทำตัวหนา+สีแดง
    ให้กับมาร์ก [ต้องระบุ: ...] (ดู ADR-002) เพื่อให้ผู้ตรวจร่างเห็นจุดที่ยังขาดข้อมูลชัดเจน
    เปิดไฟล์ Word แล้วไม่กลืนไปกับข้อความปกติ — ใช้ร่วมกันทั้งย่อหน้าปกติและเซลล์ตาราง"""
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group("bold") is not None:
            run = paragraph.add_run(m.group("bold"))
            run.bold = True
        else:
            run = paragraph.add_run(m.group("marker"))
            run.bold = True
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


_TABLE_ROW_RE = re.compile(r"^\|.*\|$")
_TABLE_SEP_CELL_RE = re.compile(r"^:?-{2,}:?$")


def _is_table_row(line: str) -> bool:
    return bool(_TABLE_ROW_RE.match(line.strip()))


def _is_table_separator_row(line: str) -> bool:
    """เช็คว่าเป็นบรรทัดคั่น header ของ markdown table หรือไม่ (เช่น '| :--- | :---: | ---|')
    ถ้าใช่ ไม่ต้องเอาไปใส่ในตารางจริง (ใช้แค่บอกว่าแถวก่อนหน้าเป็น header)"""
    stripped = line.strip()
    if not _is_table_row(stripped):
        return False
    cells = _split_table_row(stripped)
    return bool(cells) and all(_TABLE_SEP_CELL_RE.match(c) for c in cells)


def _split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [c.strip() for c in stripped.split("|")]


def _add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    """แปลงบล็อก markdown table (แถว header + แถวคั่น + แถวข้อมูล) เป็น docx table จริง —
    ก่อนหน้านี้ _add_markdown_line() ไม่รู้จัก '|...|' เลย เลยเอาทั้งบรรทัดไปใส่เป็น paragraph
    ธรรมดา ทำให้เห็น '| :--- | :--- |' โผล่มาตรงๆ ในเอกสาร (บั๊กที่แก้ในรอบนี้)
    รองรับ <br> ในเซลล์ (ตัดเป็นบรรทัดใหม่ในเซลล์เดียวกัน แทนที่จะโชว์ '<br>' ตรงๆ) และ
    **bold**/[ต้องระบุ: ...] เหมือนย่อหน้าปกติ ผ่าน _add_runs_with_bold() ตัวเดียวกัน"""
    header_line = table_lines[0]
    data_lines = [ln for ln in table_lines[1:] if not _is_table_separator_row(ln)]
    rows_raw = [header_line] + data_lines

    parsed_rows = [_split_table_row(r) for r in rows_raw]
    n_cols = max(len(r) for r in parsed_rows)
    parsed_rows = [r + [""] * (n_cols - len(r)) for r in parsed_rows]  # กันแถวเซลล์ไม่ครบ

    table = doc.add_table(rows=len(parsed_rows), cols=n_cols)
    table.style = "Table Grid"

    for ri, row_cells in enumerate(parsed_rows):
        for ci, cell_text in enumerate(row_cells):
            cell = table.cell(ri, ci)
            cell.text = ""
            para = cell.paragraphs[0]
            parts = cell_text.split("<br>")
            for pi, part in enumerate(parts):
                if pi > 0:
                    para.add_run().add_break()
                _add_runs_with_bold(para, part.strip())
            if ri == 0:
                for run in para.runs:
                    run.bold = True

    doc.add_paragraph("")  # เว้นบรรทัดหลังตาราง กันเนื้อหาถัดไปติดขอบตาราง


def markdown_to_docx(markdown_text: str, title: str) -> bytes:
    """แปลง markdown_text เป็นไฟล์ .docx (คืนค่าเป็น bytes พร้อมส่งผ่าน st.download_button)
    title จะถูกใส่เป็น Heading 0 (Title style) อยู่บนสุดของเอกสาร"""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "TH Sarabun New"
    style.font.size = Pt(14)

    if title:
        doc.add_heading(title, level=0)

    lines = markdown_text.splitlines()
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        # ตรวจจับตาราง markdown: แถวปัจจุบันเป็น '|...|' และแถวถัดไปเป็นแถวคั่น (:---) —
        # ถ้าใช่ ไล่เก็บทุกแถวที่ขึ้นต้น/ลงท้ายด้วย '|' ต่อเนื่องกันมาแปลงเป็น docx table เดียว
        if (
            _is_table_row(line)
            and i + 1 < n
            and _is_table_separator_row(lines[i + 1])
        ):
            j = i + 1
            while j < n and _is_table_row(lines[j]):
                j += 1
            _add_markdown_table(doc, lines[i:j])
            i = j
            continue
        _add_markdown_line(doc, line)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
